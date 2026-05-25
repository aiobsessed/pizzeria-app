from collections import Counter
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from io import BytesIO

from app.core.enums import OrderStatus, PaymentMethod
from app.core.timezone import MSK
from app.models import Order

_STATUS_LABELS = {
    "accepted":   "Принят",
    "preparing":  "Готовится",
    "on_the_way": "В пути",
    "delivered":  "Доставлен",
    "canceled":   "Отменён",
}
_DELIVERY_LABELS = {
    "delivery": "Доставка",
    "pickup":   "Самовывоз",
}
_PAYMENT_LABELS = {
    "cash":   "Наличные",
    "card":   "Карта",
    "online": "Онлайн",
}


@dataclass
class ReportStats:
    total_orders:    int
    delivered_count: int
    cancelled_count: int
    revenue:         Decimal
    by_status:       dict
    by_payment:      dict
    top_products:    list[tuple[str, int]]


def compute_stats(orders: list[Order]) -> ReportStats:
    delivered = [o for o in orders if o.status == OrderStatus.delivered]
    qty: Counter = Counter()
    for order in orders:
        for item in order.items:
            qty[item.product.name] += item.quantity
    return ReportStats(
        total_orders    = len(orders),
        delivered_count = len(delivered),
        cancelled_count = sum(1 for o in orders if o.status == OrderStatus.canceled),
        revenue         = sum((o.total_price for o in delivered), Decimal(0)),
        by_status       = {s: sum(1 for o in orders if o.status == s) for s in OrderStatus},
        by_payment      = {m: sum(1 for o in orders if o.payment_method == m) for m in PaymentMethod},
        top_products    = qty.most_common(10),
    )


def build_excel(orders: list[Order], date_from: date | None, date_to: date | None) -> BytesIO:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    INDIGO   = PatternFill("solid", fgColor="4F46E5")
    WHITE_FG = Font(color="FFFFFF", bold=True)
    CENTER   = Alignment(horizontal="center", vertical="center", wrap_text=True)

    wb     = openpyxl.Workbook()
    stats  = compute_stats(orders)
    period = f"{date_from or 'начало'} — {date_to or 'конец'}"

    ws = wb.active
    ws.title = "Сводка"

    ws.append(["Отчёт пиццерии", period])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])

    for label, value in [
        ("Всего заказов",               stats.total_orders),
        ("Доставлено",                  stats.delivered_count),
        ("Отменено",                    stats.cancelled_count),
        ("Выручка (доставленные), руб.", float(stats.revenue)),
    ]:
        ws.append([label, value])

    ws.append([])
    ws.append(["Статус", "Количество"])
    for cell in ws[ws.max_row]:
        cell.fill = INDIGO
        cell.font = WHITE_FG
    for status, cnt in stats.by_status.items():
        ws.append([_STATUS_LABELS[status.value], cnt])

    ws.append([])
    ws.append(["Способ оплаты", "Количество"])
    for cell in ws[ws.max_row]:
        cell.fill = INDIGO
        cell.font = WHITE_FG
    for method, cnt in stats.by_payment.items():
        ws.append([_PAYMENT_LABELS[method.value], cnt])

    ws.append([])
    ws.append(["Товар", "Продано (шт.)"])
    for cell in ws[ws.max_row]:
        cell.fill = INDIGO
        cell.font = WHITE_FG
    for name, qty in stats.top_products:
        ws.append([name, qty])

    ws.column_dimensions["A"].width = 35
    ws.column_dimensions["B"].width = 20

    ws2 = wb.create_sheet("Заказы")
    headers = ["#", "Дата", "Клиент", "Позиции", "Итого, руб.", "Статус", "Доставка", "Оплата", "Курьер"]
    ws2.append(headers)
    for cell in ws2[1]:
        cell.fill = INDIGO
        cell.font = WHITE_FG
        cell.alignment = CENTER

    for order in orders:
        items_str    = ", ".join(f"{i.product.name} x{i.quantity}" for i in order.items)
        courier_name = order.courier.name if order.courier else "—"
        ws2.append([
            order.id,
            order.created_at.astimezone(MSK).strftime("%d.%m.%Y %H:%M"),
            order.client.email,
            items_str,
            float(order.total_price),
            _STATUS_LABELS[order.status.value],
            _DELIVERY_LABELS[order.delivery_type.value],
            _PAYMENT_LABELS[order.payment_method.value],
            courier_name,
        ])

    for col, width in zip("ABCDEFGHI", [6, 18, 22, 55, 14, 14, 14, 12, 20]):
        ws2.column_dimensions[col].width = width

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _get_cyrillic_font() -> str:
    import os
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = [
        "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/freefont/FreeSans.ttf",
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "C:\\Windows\\Fonts\\arial.ttf",
    ]
    for path in candidates:
        if os.path.exists(path):
            pdfmetrics.registerFont(TTFont("CyrFont", path))
            return "CyrFont"
    return "Helvetica"


_PAGE_W = 267  # usable width on landscape A4 with 15mm margins (mm)


def build_pdf(orders: list[Order], date_from: date | None, date_to: date | None) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    INDIGO     = colors.HexColor("#4F46E5")
    INDIGO_LT  = colors.HexColor("#EEF2FF")
    GRID_COLOR = colors.HexColor("#C7D2FE")
    ROW_ALT    = colors.HexColor("#F5F3FF")
    GRAY_TEXT  = colors.HexColor("#6B7280")
    font       = _get_cyrillic_font()
    W          = _PAGE_W * mm

    def _style(name: str, **kw) -> ParagraphStyle:
        return ParagraphStyle(name, fontName=font, **kw)

    def _section_label(text: str) -> Paragraph:
        return Paragraph(text.upper(), _style(
            f"lbl_{text}", fontSize=7, textColor=GRAY_TEXT,
            spaceBefore=6, spaceAfter=3,
        ))

    def _stat_table(header: list[str], rows: list[list[str]], col_widths: list) -> Table:
        data = [header] + rows
        alt  = [("BACKGROUND", (0, i), (-1, i), ROW_ALT) for i in range(2, len(data), 2)]
        t    = Table(data, colWidths=col_widths)
        t.setStyle(TableStyle([
            ("FONTNAME",      (0, 0), (-1, -1), font),
            ("FONTSIZE",      (0, 0), (-1, -1), 8),
            ("BACKGROUND",    (0, 0), (-1, 0),  INDIGO),
            ("TEXTCOLOR",     (0, 0), (-1, 0),  colors.white),
            ("ALIGN",         (0, 0), (-1, 0),  "CENTER"),
            ("ALIGN",         (1, 1), (1,  -1), "RIGHT"),
            ("VALIGN",        (0, 0), (-1, -1), "MIDDLE"),
            ("BOX",           (0, 0), (-1, -1), 0.5,  GRID_COLOR),
            ("INNERGRID",     (0, 0), (-1, -1), 0.25, GRID_COLOR),
            ("TOPPADDING",    (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("LEFTPADDING",   (0, 0), (-1, -1), 6),
            ("RIGHTPADDING",  (0, 0), (-1, -1), 6),
        ] + alt))
        return t

    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4),
        rightMargin=15*mm, leftMargin=15*mm,
        topMargin=15*mm,   bottomMargin=15*mm,
    )
    stats  = compute_stats(orders)
    period = f"{date_from or 'начало'} — {date_to or 'конец'}"
    doc.title = f"Отчёт пиццерии {period}"

    story: list = []

    story.append(Paragraph(
        "Отчёт пиццерии",
        _style("title", fontSize=18, textColor=INDIGO, spaceAfter=11),
    ))
    story.append(Paragraph(
        f"Период: {period}",
        _style("sub", fontSize=9, textColor=GRAY_TEXT, spaceAfter=0),
    ))
    story.append(Spacer(1, 7*mm))

    kpi_w = W / 4
    t_kpi = Table(
        [
            ["Всего заказов", "Доставлено", "Отменено", "Выручка, руб."],
            [
                str(stats.total_orders),
                str(stats.delivered_count),
                str(stats.cancelled_count),
                f"{stats.revenue:,.2f}",
            ],
        ],
        colWidths=[kpi_w] * 4,
        rowHeights=[10*mm, 14*mm],
    )
    t_kpi.setStyle(TableStyle([
        ("FONTNAME",      (0, 0), (-1, -1), font),
        ("BACKGROUND",    (0, 0), (-1, 0),  INDIGO),
        ("TEXTCOLOR",     (0, 0), (-1, 0),  colors.white),
        ("FONTSIZE",      (0, 0), (-1, 0),  8),
        ("BACKGROUND",    (0, 1), (-1, 1),  INDIGO_LT),
        ("TEXTCOLOR",     (0, 1), (-1, 1),  INDIGO),
        ("FONTSIZE",      (0, 1), (-1, 1),  16),
        ("ALIGN",         (0, 0), (-1, -1), "CENTER"),
        ("VALIGN",        (0, 0), (-1, -1), "MIDDLE"),
        ("BOX",           (0, 0), (-1, -1), 0,   colors.white),
        ("INNERGRID",     (0, 0), (-1, -1), 2,   colors.white),
        ("TOPPADDING",    (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(t_kpi)
    story.append(Spacer(1, 7*mm))

    gap    = 9*mm
    pair_w = (W - gap) / 2
    lbl_w  = pair_w - 30*mm
    cnt_w  = 30*mm

    t_status  = _stat_table(
        ["Статус", "Кол-во"],
        [[_STATUS_LABELS[s.value], str(cnt)] for s, cnt in stats.by_status.items()],
        [lbl_w, cnt_w],
    )
    t_payment = _stat_table(
        ["Способ оплаты", "Кол-во"],
        [[_PAYMENT_LABELS[m.value], str(cnt)] for m, cnt in stats.by_payment.items()],
        [lbl_w, cnt_w],
    )

    story.append(_section_label("Аналитика"))
    t_pair = Table([[t_status, "", t_payment]], colWidths=[pair_w, gap, pair_w])
    t_pair.setStyle(TableStyle([
        ("LEFTPADDING",   (0, 0), (-1, -1), 0),
        ("RIGHTPADDING",  (0, 0), (-1, -1), 0),
        ("TOPPADDING",    (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
        ("VALIGN",        (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(t_pair)

    if stats.top_products:
        story.append(Spacer(1, 5*mm))
        story.append(_section_label("Топ-10 товаров"))
        story.append(_stat_table(
            ["Товар", "Продано, шт."],
            [[name, str(qty)] for name, qty in stats.top_products],
            [W - 35*mm, 35*mm],
        ))

    story.append(PageBreak())
    story.append(Paragraph(
        "Отчёт пиццерии",
        _style("title2", fontSize=14, textColor=INDIGO, spaceAfter=9),
    ))
    story.append(Paragraph(
        f"Период: {period}  |  Заказы",
        _style("sub2", fontSize=8, textColor=GRAY_TEXT, spaceAfter=0),
    ))
    story.append(Spacer(1, 4*mm))

    # 10 + 24 + 46 + 97 + 22 + 22 + 22 + 24 = 267mm = W
    col_widths = [10*mm, 24*mm, 46*mm, 97*mm, 22*mm, 22*mm, 22*mm, 24*mm]
    cell_style = _style("cell", fontSize=8, leading=10)
    rows = [["#", "Дата", "Клиент", "Позиции", "Итого, руб.", "Статус", "Доставка", "Оплата"]]
    for o in orders:
        items_str = ", ".join(f"{i.product.name} x{i.quantity}" for i in o.items)
        rows.append([
            str(o.id),
            o.created_at.astimezone(MSK).strftime("%d.%m.%Y\n%H:%M"),
            Paragraph(o.client.email, cell_style),
            Paragraph(items_str, cell_style),
            f"{o.total_price:,.2f}",
            _STATUS_LABELS[o.status.value],
            _DELIVERY_LABELS[o.delivery_type.value],
            _PAYMENT_LABELS[o.payment_method.value],
        ])

    alt_cmds = [("BACKGROUND", (0, i), (-1, i), ROW_ALT) for i in range(2, len(rows), 2)]
    t_orders = Table(rows, colWidths=col_widths, repeatRows=1)
    t_orders.setStyle(TableStyle([
        ("FONTNAME",      (0, 0), (-1, -1), font),
        ("FONTSIZE",      (0, 0), (-1, -1), 8),
        ("BACKGROUND",    (0, 0), (-1, 0),  INDIGO),
        ("TEXTCOLOR",     (0, 0), (-1, 0),  colors.white),
        ("ALIGN",         (0, 0), (-1, 0),  "CENTER"),
        ("ALIGN",         (4, 1), (4,  -1), "RIGHT"),
        ("VALIGN",        (0, 0), (-1, -1), "TOP"),
        ("BOX",           (0, 0), (-1, -1), 0.5,  GRID_COLOR),
        ("INNERGRID",     (0, 0), (-1, -1), 0.25, GRID_COLOR),
        ("TOPPADDING",    (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING",   (0, 0), (-1, -1), 4),
        ("RIGHTPADDING",  (0, 0), (-1, -1), 4),
    ] + alt_cmds))
    story.append(t_orders)

    doc.build(story)
    return buf.getvalue()
