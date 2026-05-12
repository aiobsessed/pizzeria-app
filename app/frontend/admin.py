from collections import Counter
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from io import BytesIO

from fastapi import APIRouter, Depends, Form, Request
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.dependencies import flash_redirect, get_db, get_flash, require_admin_from_cookie
from app.core.enums import DeliveryType, EmployeeStatus, OrderStatus, PaymentMethod
from app.core.exceptions import ConflictError, NotFoundError
from app.models import Employee, Order
from app.schemas import (
    CategoryCreate,
    CategoryUpdate,
    EmployeeCreate,
    EmployeeUpdate,
    OrderUpdate,
    ProductCreate,
    ProductUpdate,
)
from app.services import (
    PositionService,
    CategoryService,
    ClientService,
    EmployeeService,
    OrderService,
    ProductService,
)

router = APIRouter(prefix="/admin", tags=["frontend-admin"])
templates = Jinja2Templates(directory="app/templates")


# ── Report helpers ─────────────────────────────────────────────────────────────

_STATUS_LABELS   = {"accepted": "Принят", "preparing": "Готовится", "on_the_way": "В пути", "delivered": "Доставлен", "canceled": "Отменён"}
_DELIVERY_LABELS = {"delivery": "Доставка", "pickup": "Самовывоз"}
_PAYMENT_LABELS  = {"cash": "Наличные", "card": "Карта", "online": "Онлайн"}


@dataclass
class _ReportStats:
    total_orders:    int
    delivered_count: int
    cancelled_count: int
    revenue:         Decimal
    by_status:       dict
    by_payment:      dict
    top_products:    list[tuple[str, int]]


def _filter_orders(orders: list[Order], date_from: date | None, date_to: date | None) -> list[Order]:
    if date_from:
        orders = [o for o in orders if o.created_at.date() >= date_from]
    if date_to:
        orders = [o for o in orders if o.created_at.date() <= date_to]
    return orders


def _compute_stats(orders: list[Order]) -> _ReportStats:
    delivered = [o for o in orders if o.status == OrderStatus.delivered]
    qty: Counter = Counter()
    for order in orders:
        for item in order.items:
            qty[item.product.name] += item.quantity
    return _ReportStats(
        total_orders    = len(orders),
        delivered_count = len(delivered),
        cancelled_count = sum(1 for o in orders if o.status == OrderStatus.canceled),
        revenue         = sum((o.total_price for o in delivered), Decimal(0)),
        by_status       = {s: sum(1 for o in orders if o.status == s) for s in OrderStatus},
        by_payment      = {m: sum(1 for o in orders if o.payment_method == m) for m in PaymentMethod},
        top_products    = qty.most_common(10),
    )


def _build_excel(orders: list[Order], date_from: date | None, date_to: date | None) -> BytesIO:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    INDIGO   = PatternFill("solid", fgColor="4F46E5")
    WHITE_FG = Font(color="FFFFFF", bold=True)
    CENTER   = Alignment(horizontal="center", vertical="center", wrap_text=True)

    wb = openpyxl.Workbook()
    stats = _compute_stats(orders)
    period = f"{date_from or 'начало'} — {date_to or 'конец'}"

    ws = wb.active
    ws.title = "Сводка"

    ws.append(["Отчёт пиццерии", period])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])

    for label, value in [
        ("Всего заказов",             stats.total_orders),
        ("Доставлено",                stats.delivered_count),
        ("Отменено",                  stats.cancelled_count),
        ("Выручка (доставленные), ₽", float(stats.revenue)),
    ]:
        ws.append([label, value])

    ws.append([])
    ws.append(["Статус", "Количество"])
    for cell in ws[ws.max_row]:
        cell.fill = INDIGO
        cell.font = WHITE_FG
    for status, cnt in stats.by_status.items():
        ws.append([_STATUS_LABELS[status.value], cnt])

    ws.append([])
    ws.append(["Способ оплаты", "Количество"])
    for cell in ws[ws.max_row]:
        cell.fill = INDIGO
        cell.font = WHITE_FG
    for method, cnt in stats.by_payment.items():
        ws.append([_PAYMENT_LABELS[method.value], cnt])

    ws.append([])
    ws.append(["Товар", "Продано (шт.)"])
    for cell in ws[ws.max_row]:
        cell.fill = INDIGO
        cell.font = WHITE_FG
    for name, qty in stats.top_products:
        ws.append([name, qty])

    ws.column_dimensions["A"].width = 35
    ws.column_dimensions["B"].width = 20

    ws2 = wb.create_sheet("Заказы")
    headers = ["#", "Дата", "Клиент", "Позиции", "Итого, ₽", "Статус", "Доставка", "Оплата", "Курьер"]
    ws2.append(headers)
    for cell in ws2[1]:
        cell.fill = INDIGO
        cell.font = WHITE_FG
        cell.alignment = CENTER

    for order in orders:
        items_str = ", ".join(f"{i.product.name} ×{i.quantity}" for i in order.items)
        courier_name = order.courier.name if order.courier else "—"
        ws2.append([
            order.id,
            order.created_at.strftime("%d.%m.%Y %H:%M"),
            order.client.name,
            items_str,
            float(order.total_price),
            _STATUS_LABELS[order.status.value],
            _DELIVERY_LABELS[order.delivery_type.value],
            _PAYMENT_LABELS[order.payment_method.value],
            courier_name,
        ])

    for col, width in zip("ABCDEFGHI", [6, 18, 22, 55, 14, 14, 14, 12, 20]):
        ws2.column_dimensions[col].width = width

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _get_cyrillic_font() -> str:
    import os
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = [
        "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/freefont/FreeSans.ttf",
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "C:\\Windows\\Fonts\\arial.ttf",
    ]
    for path in candidates:
        if os.path.exists(path):
            pdfmetrics.registerFont(TTFont("CyrFont", path))
            return "CyrFont"
    return "Helvetica"


def _build_pdf(orders: list[Order], date_from: date | None, date_to: date | None) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    INDIGO     = colors.HexColor("#4F46E5")
    INDIGO_LT  = colors.HexColor("#EEF2FF")
    GRID_COLOR = colors.HexColor("#C7D2FE")
    ROW_ALT    = colors.HexColor("#F5F3FF")
    font       = _get_cyrillic_font()

    def _style(name: str, **kw) -> ParagraphStyle:
        return ParagraphStyle(name, fontName=font, **kw)

    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                            rightMargin=15*mm, leftMargin=15*mm,
                            topMargin=15*mm, bottomMargin=15*mm)

    stats  = _compute_stats(orders)
    period = f"{date_from or 'начало'} — {date_to or 'конец'}"

    story: list = []
    story.append(Paragraph(f"Отчёт пиццерии  |  Период: {period}",
                            _style("title", fontSize=15, spaceAfter=5)))
    story.append(Spacer(1, 5*mm))

    summary_data = [
        ["Всего заказов", "Доставлено", "Отменено", "Выручка, ₽"],
        [str(stats.total_orders), str(stats.delivered_count),
         str(stats.cancelled_count), f"{stats.revenue:.2f}"],
    ]
    w = 62*mm
    t_summary = Table(summary_data, colWidths=[w, w, w, w])
    t_summary.setStyle(TableStyle([
        ("FONTNAME",      (0, 0), (-1, -1), font),
        ("BACKGROUND",    (0, 0), (-1, 0),  INDIGO),
        ("TEXTCOLOR",     (0, 0), (-1, 0),  colors.white),
        ("BACKGROUND",    (0, 1), (-1, 1),  INDIGO_LT),
        ("FONTSIZE",      (0, 0), (-1, 0),  9),
        ("FONTSIZE",      (0, 1), (-1, 1),  14),
        ("ALIGN",         (0, 0), (-1, -1), "CENTER"),
        ("VALIGN",        (0, 0), (-1, -1), "MIDDLE"),
        ("BOX",           (0, 0), (-1, -1), 0.5, GRID_COLOR),
        ("INNERGRID",     (0, 0), (-1, -1), 0.5, GRID_COLOR),
        ("TOPPADDING",    (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(t_summary)
    story.append(Spacer(1, 7*mm))
    story.append(Paragraph("Заказы", _style("h2", fontSize=11, spaceBefore=4, spaceAfter=4)))

    col_widths = [12*mm, 26*mm, 33*mm, 88*mm, 22*mm, 22*mm, 22*mm, 22*mm]
    rows = [["#", "Дата", "Клиент", "Позиции", "Итого, ₽", "Статус", "Доставка", "Оплата"]]
    for o in orders:
        items_str = ", ".join(f"{i.product.name} ×{i.quantity}" for i in o.items)
        rows.append([
            str(o.id),
            o.created_at.strftime("%d.%m.%Y\n%H:%M"),
            o.client.name,
            items_str,
            f"{o.total_price:.2f}",
            _STATUS_LABELS[o.status.value],
            _DELIVERY_LABELS[o.delivery_type.value],
            _PAYMENT_LABELS[o.payment_method.value],
        ])

    base_cmds = [
        ("FONTNAME",      (0, 0), (-1, -1), font),
        ("FONTSIZE",      (0, 0), (-1, -1), 8),
        ("BACKGROUND",    (0, 0), (-1, 0),  INDIGO),
        ("TEXTCOLOR",     (0, 0), (-1, 0),  colors.white),
        ("ALIGN",         (0, 0), (-1, 0),  "CENTER"),
        ("ALIGN",         (4, 1), (4, -1),  "RIGHT"),
        ("VALIGN",        (0, 0), (-1, -1), "TOP"),
        ("BOX",           (0, 0), (-1, -1), 0.5, GRID_COLOR),
        ("INNERGRID",     (0, 0), (-1, -1), 0.25, GRID_COLOR),
        ("TOPPADDING",    (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    alt_cmds = [
        ("BACKGROUND", (0, i), (-1, i), ROW_ALT)
        for i in range(2, len(rows), 2)
    ]

    t_orders = Table(rows, colWidths=col_widths, repeatRows=1)
    t_orders.setStyle(TableStyle(base_cmds + alt_cmds))
    story.append(t_orders)

    doc.build(story)
    return buf.getvalue()


# ── Dashboard ─────────────────────────────────────────────────────────────────


@router.get("", response_class=HTMLResponse)
async def dashboard(
    request: Request,
    employee: Employee = Depends(require_admin_from_cookie),
    session: AsyncSession = Depends(get_db),
    flash: str | None = Depends(get_flash),
) -> HTMLResponse:
    orders   = await OrderService(session).get_all_with_items()
    clients  = await ClientService(session).get_all()
    products = await ProductService(session).get_all()
    couriers = await EmployeeService(session).get_all(role="courier")

    response = templates.TemplateResponse(
        request,
        "admin/dashboard.html",
        {
            "employee":         employee,
            "flash":            flash,
            "orders_count":     len(orders),
            "clients_count":    len(clients),
            "products_count":   len(products),
            "couriers_count":   len(couriers),
            "orders_by_status": {s: sum(1 for o in orders if o.status == s) for s in OrderStatus},
            "OrderStatus":      OrderStatus,
        },
    )
    response.delete_cookie("flash")
    return response


# ── Orders ────────────────────────────────────────────────────────────────────


@router.get("/orders", response_class=HTMLResponse)
async def orders_page(
    request: Request,
    status: str | None = None,
    employee: Employee = Depends(require_admin_from_cookie),
    session: AsyncSession = Depends(get_db),
    flash: str | None = Depends(get_flash),
) -> HTMLResponse:
    status_enum: OrderStatus | None = None
    if status:
        try:
            status_enum = OrderStatus(status)
        except ValueError:
            pass

    orders   = await OrderService(session).get_all_with_items(status=status_enum)
    couriers = await EmployeeService(session).get_all(role="courier")

    response = templates.TemplateResponse(
        request,
        "admin/orders.html",
        {
            "employee":       employee,
            "flash":          flash,
            "orders":         orders,
            "couriers":       couriers,
            "OrderStatus":    OrderStatus,
            "DeliveryType":   DeliveryType,
            "PaymentMethod":  PaymentMethod,
            "current_status": status_enum,
        },
    )
    response.delete_cookie("flash")
    return response


@router.patch("/orders/{order_id}/status", response_class=HTMLResponse)
async def update_order_status(
    request: Request,
    order_id: int,
    status: OrderStatus = Form(),
    _: Employee = Depends(require_admin_from_cookie),
    session: AsyncSession = Depends(get_db),
) -> HTMLResponse:
    try:
        order = await OrderService(session).update(order_id, OrderUpdate(status=status))
    except NotFoundError:
        return HTMLResponse("", status_code=404)

    couriers = await EmployeeService(session).get_all(role="courier")
    return templates.TemplateResponse(
        request,
        "admin/partials/order_row.html",
        {"order": order, "couriers": couriers, "OrderStatus": OrderStatus},
    )


@router.patch("/orders/{order_id}/courier", response_class=HTMLResponse)
async def assign_courier(
    request: Request,
    order_id: int,
    courier_id: str = Form(default=""),
    _: Employee = Depends(require_admin_from_cookie),
    session: AsyncSession = Depends(get_db),
) -> HTMLResponse:
    courier_id_int: int | None = int(courier_id) if courier_id.strip() else None
    try:
        order = await OrderService(session).update(order_id, OrderUpdate(courier_id=courier_id_int))
    except NotFoundError:
        return HTMLResponse("", status_code=404)

    couriers = await EmployeeService(session).get_all(role="courier")
    return templates.TemplateResponse(
        request,
        "admin/partials/order_row.html",
        {"order": order, "couriers": couriers, "OrderStatus": OrderStatus},
    )


# ── Clients ───────────────────────────────────────────────────────────────────


@router.get("/clients", response_class=HTMLResponse)
async def clients_page(
    request: Request,
    employee: Employee = Depends(require_admin_from_cookie),
    session: AsyncSession = Depends(get_db),
    flash: str | None = Depends(get_flash),
) -> HTMLResponse:
    clients = await ClientService(session).get_all()

    response = templates.TemplateResponse(
        request,
        "admin/clients.html",
        {"employee": employee, "flash": flash, "clients": clients},
    )
    response.delete_cookie("flash")
    return response


@router.patch("/clients/{client_id}/block", response_class=HTMLResponse)
async def block_client(
    request: Request,
    client_id: int,
    _: Employee = Depends(require_admin_from_cookie),
    session: AsyncSession = Depends(get_db),
) -> HTMLResponse:
    try:
        client = await ClientService(session).block(client_id)
    except NotFoundError:
        return HTMLResponse("", status_code=404)

    return templates.TemplateResponse(
        request,
        "admin/partials/client_row.html",
        {"client": client},
    )


# ── Products ──────────────────────────────────────────────────────────────────


@router.get("/products", response_class=HTMLResponse)
async def products_page(
    request: Request,
    employee: Employee = Depends(require_admin_from_cookie),
    session: AsyncSession = Depends(get_db),
    flash: str | None = Depends(get_flash),
) -> HTMLResponse:
    products   = await ProductService(session).get_all()
    categories = await CategoryService(session).get_all()

    response = templates.TemplateResponse(
        request,
        "admin/products.html",
        {
            "employee":   employee,
            "flash":      flash,
            "products":   products,
            "categories": categories,
        },
    )
    response.delete_cookie("flash")
    return response


@router.post("/products/create")
async def create_product(
    name: str         = Form(),
    category_id: int  = Form(),
    weight: int       = Form(),
    price: str        = Form(),
    description: str  = Form(default=""),
    composition: str  = Form(default=""),
    image_url: str    = Form(default=""),
    is_available: str = Form(default=""),
    _: Employee       = Depends(require_admin_from_cookie),
    session: AsyncSession = Depends(get_db),
) -> RedirectResponse:
    try:
        await ProductService(session).create(
            ProductCreate(
                name=name,
                category_id=category_id,
                weight=weight,
                price=Decimal(price),
                description=description.strip() or None,
                composition=composition.strip() or None,
                image_url=image_url.strip() or None,
                is_available=is_available == "on",
            )
        )
    except (NotFoundError, ConflictError) as e:
        return flash_redirect("/admin/products", str(e))

    return flash_redirect("/admin/products", "Товар создан", success=True)


@router.post("/products/{product_id}/update")
async def update_product(
    product_id: int,
    name: str         = Form(),
    category_id: int  = Form(),
    weight: int       = Form(),
    price: str        = Form(),
    description: str  = Form(default=""),
    composition: str  = Form(default=""),
    image_url: str    = Form(default=""),
    is_available: str = Form(default=""),
    _: Employee       = Depends(require_admin_from_cookie),
    session: AsyncSession = Depends(get_db),
) -> RedirectResponse:
    try:
        await ProductService(session).update(
            product_id,
            ProductUpdate(
                name=name,
                category_id=category_id,
                weight=weight,
                price=Decimal(price),
                description=description.strip() or None,
                composition=composition.strip() or None,
                image_url=image_url.strip() or None,
                is_available=is_available == "on",
            ),
        )
    except (NotFoundError, ConflictError) as e:
        return flash_redirect("/admin/products", str(e))

    return flash_redirect("/admin/products", "Товар обновлён", success=True)


@router.post("/products/{product_id}/delete")
async def delete_product(
    product_id: int,
    _: Employee = Depends(require_admin_from_cookie),
    session: AsyncSession = Depends(get_db),
) -> RedirectResponse:
    try:
        await ProductService(session).delete(product_id)
    except (NotFoundError, ConflictError) as e:
        return flash_redirect("/admin/products", str(e))

    return flash_redirect("/admin/products", "Товар деактивирован", success=True)


# ── Categories ────────────────────────────────────────────────────────────────


@router.get("/categories", response_class=HTMLResponse)
async def categories_page(
    request: Request,
    employee: Employee = Depends(require_admin_from_cookie),
    session: AsyncSession = Depends(get_db),
    flash: str | None = Depends(get_flash),
) -> HTMLResponse:
    categories = await CategoryService(session).get_all()

    response = templates.TemplateResponse(
        request,
        "admin/categories.html",
        {"employee": employee, "flash": flash, "categories": categories},
    )
    response.delete_cookie("flash")
    return response


@router.post("/categories/create")
async def create_category(
    name: str      = Form(),
    slug: str      = Form(),
    is_active: str = Form(default=""),
    _: Employee    = Depends(require_admin_from_cookie),
    session: AsyncSession = Depends(get_db),
) -> RedirectResponse:
    try:
        await CategoryService(session).create(
            CategoryCreate(name=name, slug=slug, is_active=is_active == "on")
        )
    except ConflictError as e:
        return flash_redirect("/admin/categories", str(e))

    return flash_redirect("/admin/categories", "Категория создана", success=True)


@router.post("/categories/{category_id}/update")
async def update_category(
    category_id: int,
    name: str      = Form(),
    slug: str      = Form(),
    is_active: str = Form(default=""),
    _: Employee    = Depends(require_admin_from_cookie),
    session: AsyncSession = Depends(get_db),
) -> RedirectResponse:
    try:
        await CategoryService(session).update(
            category_id,
            CategoryUpdate(name=name, slug=slug, is_active=is_active == "on"),
        )
    except (NotFoundError, ConflictError) as e:
        return flash_redirect("/admin/categories", str(e))

    return flash_redirect("/admin/categories", "Категория обновлена", success=True)


@router.post("/categories/{category_id}/delete")
async def delete_category(
    category_id: int,
    _: Employee = Depends(require_admin_from_cookie),
    session: AsyncSession = Depends(get_db),
) -> RedirectResponse:
    try:
        await CategoryService(session).delete(category_id)
    except (NotFoundError, ConflictError) as e:
        return flash_redirect("/admin/categories", str(e))

    return flash_redirect("/admin/categories", "Категория деактивирована", success=True)


# ── Employees ─────────────────────────────────────────────────────────────────


@router.get("/employees", response_class=HTMLResponse)
async def employees_page(
    request: Request,
    employee: Employee = Depends(require_admin_from_cookie),
    session: AsyncSession = Depends(get_db),
    flash: str | None = Depends(get_flash),
) -> HTMLResponse:
    employees = await EmployeeService(session).get_all()
    positions = await PositionService(session).get_all()

    response = templates.TemplateResponse(
        request,
        "admin/employees.html",
        {
            "employee":       employee,
            "flash":          flash,
            "employees":      employees,
            "positions":      positions,
            "EmployeeStatus": EmployeeStatus,
        },
    )
    response.delete_cookie("flash")
    return response


@router.post("/employees/create")
async def create_employee(
    position_id: int = Form(),
    name: str        = Form(),
    email: str       = Form(),
    phone: str       = Form(),
    inn: str         = Form(),
    password: str    = Form(),
    _: Employee      = Depends(require_admin_from_cookie),
    session: AsyncSession = Depends(get_db),
) -> RedirectResponse:
    try:
        await EmployeeService(session).create(
            EmployeeCreate(
                position_id=position_id,
                name=name,
                email=email,
                phone=phone,
                inn=inn,
                password=password,
            )
        )
    except (NotFoundError, ConflictError) as e:
        return flash_redirect("/admin/employees", str(e))

    return flash_redirect("/admin/employees", "Сотрудник создан", success=True)


@router.post("/employees/{employee_id}/update")
async def update_employee(
    employee_id: int,
    position_id: int = Form(),
    name: str        = Form(),
    email: str       = Form(),
    phone: str       = Form(),
    inn: str         = Form(),
    status: str      = Form(default=""),
    _: Employee      = Depends(require_admin_from_cookie),
    session: AsyncSession = Depends(get_db),
) -> RedirectResponse:
    try:
        await EmployeeService(session).update(
            employee_id,
            EmployeeUpdate(
                position_id=position_id,
                name=name,
                email=email,
                phone=phone,
                inn=inn,
                status=EmployeeStatus(status) if status else None,
            ),
        )
    except (NotFoundError, ConflictError) as e:
        return flash_redirect("/admin/employees", str(e))

    return flash_redirect("/admin/employees", "Сотрудник обновлён", success=True)


@router.post("/employees/{employee_id}/delete")
async def delete_employee(
    employee_id: int,
    _: Employee = Depends(require_admin_from_cookie),
    session: AsyncSession = Depends(get_db),
) -> RedirectResponse:
    try:
        await EmployeeService(session).delete(employee_id)
    except NotFoundError as e:
        return flash_redirect("/admin/employees", str(e))

    return flash_redirect("/admin/employees", "Сотрудник удалён", success=True)


# ── Reports ───────────────────────────────────────────────────────────────────


@router.get("/reports", response_class=HTMLResponse)
async def reports_page(
    request: Request,
    date_from: date | None = None,
    date_to: date | None = None,
    employee: Employee = Depends(require_admin_from_cookie),
    session: AsyncSession = Depends(get_db),
    flash: str | None = Depends(get_flash),
) -> HTMLResponse:
    all_orders = await OrderService(session).get_all_with_items()
    orders     = _filter_orders(all_orders, date_from, date_to)
    stats      = _compute_stats(orders)

    parts = []
    if date_from:
        parts.append(f"date_from={date_from}")
    if date_to:
        parts.append(f"date_to={date_to}")
    export_qs = "?" + "&".join(parts) if parts else ""

    response = templates.TemplateResponse(
        request,
        "admin/reports.html",
        {
            "employee":      employee,
            "flash":         flash,
            "stats":         stats,
            "date_from":     date_from,
            "date_to":       date_to,
            "export_qs":     export_qs,
            "OrderStatus":   OrderStatus,
            "PaymentMethod": PaymentMethod,
        },
    )
    response.delete_cookie("flash")
    return response


@router.get("/reports/export/excel")
async def export_excel(
    date_from: date | None = None,
    date_to: date | None = None,
    _: Employee = Depends(require_admin_from_cookie),
    session: AsyncSession = Depends(get_db),
) -> StreamingResponse:
    all_orders = await OrderService(session).get_all_with_items()
    orders     = _filter_orders(all_orders, date_from, date_to)
    buf        = _build_excel(orders, date_from, date_to)
    filename   = f"report_{date_from or 'all'}_{date_to or 'all'}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/reports/export/pdf")
async def export_pdf(
    date_from: date | None = None,
    date_to: date | None = None,
    _: Employee = Depends(require_admin_from_cookie),
    session: AsyncSession = Depends(get_db),
) -> StreamingResponse:
    all_orders = await OrderService(session).get_all_with_items()
    orders     = _filter_orders(all_orders, date_from, date_to)
    pdf_bytes  = _build_pdf(orders, date_from, date_to)
    filename   = f"report_{date_from or 'all'}_{date_to or 'all'}.pdf"
    return StreamingResponse(
        BytesIO(pdf_bytes),
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
